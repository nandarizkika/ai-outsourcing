import io
import pytest
from pptx import Presentation

from src.agents.deck_agent import DeckAgent
from src.core.models import AgentResult


def test_deck_agent_returns_pptx_bytes():
    agent = DeckAgent()
    result = agent.run(
        title="Q1 Revenue Analysis",
        sections=[
            {"heading": "Key Findings", "body": "Revenue grew 15% in Q1."},
            {"heading": "Recommendations", "body": "Focus on top-performing regions."},
        ],
        charts=[],
    )
    assert result.success is True
    assert result.deck_pptx is not None
    assert len(result.deck_pptx) > 0
    prs = Presentation(io.BytesIO(result.deck_pptx))
    assert len(prs.slides) >= 1


def test_deck_agent_title_slide_text():
    agent = DeckAgent()
    result = agent.run(
        title="Monthly Report",
        sections=[{"heading": "Summary", "body": "All metrics are on track."}],
        charts=[],
    )
    assert result.success is True
    prs = Presentation(io.BytesIO(result.deck_pptx))
    first_slide_text = " ".join(
        shape.text for shape in prs.slides[0].shapes if shape.has_text_frame
    )
    assert "Monthly Report" in first_slide_text


def test_deck_agent_attaches_chart_image():
    agent = DeckAgent()
    # Minimal valid 1x1 PNG
    png_bytes = (
        b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01"
        b"\x00\x00\x00\x01\x08\x02\x00\x00\x00\x90wS\xde\x00\x00"
        b"\x00\x0cIDATx\x9cc\xf8\x0f\x00\x00\x01\x01\x00\x05\x18"
        b"\xd8N\x00\x00\x00\x00IEND\xaeB`\x82"
    )
    result = agent.run(
        title="Chart Report",
        sections=[{"heading": "Chart", "body": "See the chart below."}],
        charts=[png_bytes],
    )
    assert result.success is True
    prs = Presentation(io.BytesIO(result.deck_pptx))
    has_picture = any(
        shape.shape_type == 13
        for slide in prs.slides
        for shape in slide.shapes
    )
    assert has_picture


import json
from unittest.mock import MagicMock
from src.core.models import Storyline, FindingSlide, Solution


def test_storyline_model_fields():
    s = Storyline(
        problem_statement="Churn increased 20% in Q2.",
        executive_summary=["Churn up 20%", "Root cause: pricing"],
        key_findings=[
            FindingSlide(heading="Q2 Churn", body="Churn hit 15%.",
                         so_what="We risk losing top cohort.", chart_index=None)
        ],
        solutions=[
            Solution(title="Price rollback", description="Revert to Q1 pricing.",
                     pros=["Quick"], cons=["Revenue impact"])
        ],
        recommended_solution="Price rollback",
        recommendation_rationale="Fastest path to churn reduction.",
        conclusion="Act within 30 days.",
        next_steps=["Rollback pricing", "Monitor churn weekly"],
    )
    assert s.problem_statement == "Churn increased 20% in Q2."
    assert len(s.key_findings) == 1
    assert s.key_findings[0].chart_index is None
    assert len(s.solutions) == 1


def test_agent_result_storyline_defaults_none():
    r = AgentResult(agent_name="deck_agent", success=True)
    assert r.storyline is None


def test_build_storyline_returns_storyline_object():
    mock_llm = MagicMock()
    mock_llm.complete.return_value = json.dumps({
        "problem_statement": "Revenue dropped.",
        "executive_summary": ["Revenue down 10%"],
        "key_findings": [{"heading": "Drop", "body": "Revenue fell.", "so_what": "Action needed.", "chart_index": None}],
        "solutions": [{"title": "Fix X", "description": "Do X.", "pros": ["fast"], "cons": ["costly"]}],
        "recommended_solution": "Fix X",
        "recommendation_rationale": "Fastest fix.",
        "conclusion": "Act now.",
        "next_steps": ["Start fix"],
    })
    agent = DeckAgent(llm=mock_llm)
    storyline = agent.build_storyline(
        analysis_text="Revenue dropped 10% in Q3.",
        findings=["Revenue fell"],
        solutions=[{"title": "Fix X", "description": "Do X.", "pros": ["fast"], "cons": ["costly"]}],
        recommendation="Fix X",
    )
    assert isinstance(storyline, Storyline)
    assert storyline.problem_statement == "Revenue dropped."
    assert len(storyline.key_findings) == 1
