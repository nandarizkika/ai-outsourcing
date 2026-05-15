import io
import json
import logging
from typing import Optional

from pptx import Presentation
from pptx.util import Inches, Pt

from src.core.models import AgentResult, FindingSlide, Solution, Storyline, TaskType

_logger = logging.getLogger(__name__)


class DeckAgent:
    def __init__(self, llm=None) -> None:
        self._llm = llm

    def run(
        self,
        title: str,
        sections: list[dict],
        charts: list[bytes],
        template_path: Optional[str] = None,
    ) -> AgentResult:
        try:
            prs = Presentation(template_path) if template_path else Presentation()
            self._add_title_slide(prs, title)
            for section in sections:
                self._add_content_slide(prs, section["heading"], section["body"])
            for i, png in enumerate(charts):
                self._add_chart_slide(prs, png, f"Chart {i + 1}")
            buf = io.BytesIO()
            prs.save(buf)
            buf.seek(0)
            return AgentResult(
                agent_name="deck_agent",
                success=True,
                deck_pptx=buf.read(),
            )
        except Exception as exc:
            _logger.error("DeckAgent failed: %s", exc, exc_info=True)
            return AgentResult(agent_name="deck_agent", success=False, error=str(exc))

    def _add_title_slide(self, prs: Presentation, title: str) -> None:
        layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title:
            slide.shapes.title.text = title
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = "AI Data Analyst"

    def _add_content_slide(self, prs: Presentation, heading: str, body: str) -> None:
        layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(layout)
        if slide.shapes.title:
            slide.shapes.title.text = heading
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = body

    def _add_chart_slide(self, prs: Presentation, png_bytes: bytes, caption: str) -> None:
        layout = prs.slide_layouts[6]  # blank
        slide = prs.slides.add_slide(layout)
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.6))
        txBox.text_frame.text = caption
        slide.shapes.add_picture(
            io.BytesIO(png_bytes),
            Inches(0.5),
            Inches(1.0),
            Inches(9),
            Inches(5.5),
        )

    def build_storyline(
        self,
        analysis_text: str,
        findings: list[str],
        solutions: list[dict],
        recommendation: str,
    ) -> Storyline:
        if self._llm is None:
            return self._fallback_storyline(analysis_text, findings, solutions, recommendation)
        system = (
            "You are a consulting analyst. Structure the following analysis as a presentation narrative. "
            "Return ONLY valid JSON matching this schema exactly:\n"
            '{"problem_statement": "str", "executive_summary": ["str"], '
            '"key_findings": [{"heading": "str", "body": "str", "so_what": "str", "chart_index": null}], '
            '"solutions": [{"title": "str", "description": "str", "pros": ["str"], "cons": ["str"]}], '
            '"recommended_solution": "str", "recommendation_rationale": "str", '
            '"conclusion": "str", "next_steps": ["str"]}'
        )
        user = (
            f"Analysis:\n{analysis_text}\n\n"
            f"Key findings:\n{chr(10).join(findings)}\n\n"
            f"Solutions considered:\n{chr(10).join(s.get('title', '') for s in solutions)}\n\n"
            f"Recommendation: {recommendation}"
        )
        raw = self._llm.complete(TaskType.REASONING, system, user)
        try:
            data = json.loads(raw)
            return Storyline(
                problem_statement=data.get("problem_statement", ""),
                executive_summary=data.get("executive_summary", []),
                key_findings=[FindingSlide(**f) for f in data.get("key_findings", [])],
                solutions=[Solution(**s) for s in data.get("solutions", [])],
                recommended_solution=data.get("recommended_solution", recommendation),
                recommendation_rationale=data.get("recommendation_rationale", ""),
                conclusion=data.get("conclusion", ""),
                next_steps=data.get("next_steps", []),
            )
        except (json.JSONDecodeError, TypeError, KeyError):
            return self._fallback_storyline(analysis_text, findings, solutions, recommendation)

    def _fallback_storyline(
        self,
        analysis_text: str,
        findings: list[str],
        solutions: list[dict],
        recommendation: str,
    ) -> Storyline:
        paragraphs = [p.strip() for p in analysis_text.split("\n\n") if p.strip()]
        return Storyline(
            problem_statement=paragraphs[0] if paragraphs else analysis_text[:200],
            executive_summary=findings[:5],
            key_findings=[
                FindingSlide(heading=f"Finding {i+1}", body=p, so_what="")
                for i, p in enumerate(paragraphs[:5])
            ],
            solutions=[
                Solution(title=s.get("title", f"Option {i+1}"),
                         description=s.get("description", ""),
                         pros=s.get("pros", []), cons=s.get("cons", []))
                for i, s in enumerate(solutions)
            ],
            recommended_solution=recommendation,
            recommendation_rationale="",
            conclusion="",
            next_steps=[],
        )
